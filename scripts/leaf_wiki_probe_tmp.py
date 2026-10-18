#!/usr/bin/env python3
"""Temporary, throttled official-Warframe-Wiki leaf probe for the acquisition audit."""

from __future__ import annotations

import json
import threading
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "acquisition-without-source.xlsx"
OUT = Path("/tmp/wiki-acq-search/leaf-wiki-hits.json")
API = "https://wiki.warframe.com/api.php"
WORKERS = 3
MIN_REQUEST_INTERVAL = 0.20

_rate_lock = threading.Lock()
_next_request_at = 0.0


def get_json(params: dict[str, str | int]) -> dict:
    global _next_request_at
    with _rate_lock:
        now = time.monotonic()
        wait = max(0.0, _next_request_at - now)
        if wait:
            time.sleep(wait)
        _next_request_at = time.monotonic() + MIN_REQUEST_INTERVAL
    url = API + "?" + urllib.parse.urlencode(params)
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Kiedas-Orbiter acquisition source audit/2026-10-16"},
    )
    last_error = None
    for attempt in range(4):
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                return json.load(response)
        except Exception as error:  # network/API errors are recorded after bounded retries
            last_error = error
            time.sleep(1.0 * (attempt + 1))
    raise RuntimeError(f"{type(last_error).__name__}: {last_error}")


def probe_leaf(leaf: str) -> dict:
    result = {"leaf": leaf, "allimages": [], "search": [], "errors": []}
    try:
        payload = get_json(
            {
                "action": "query",
                "list": "allimages",
                "aiprefix": leaf,
                "ailimit": 10,
                "format": "json",
                "formatversion": 2,
            }
        )
        result["allimages"] = payload.get("query", {}).get("allimages", [])
    except Exception as error:
        result["errors"].append({"query": "allimages", "error": str(error)})
    try:
        payload = get_json(
            {
                "action": "query",
                "list": "search",
                "srsearch": leaf,
                "srnamespace": "0|6",
                "srlimit": 10,
                "format": "json",
                "formatversion": 2,
            }
        )
        result["search"] = payload.get("query", {}).get("search", [])
    except Exception as error:
        result["errors"].append({"query": "search", "error": str(error)})
    return result


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    sheet = workbook["No source"]
    rows = []
    for row_number in range(2, sheet.max_row + 1):
        place = sheet.cell(row_number, 13).value
        if place not in (None, ""):
            continue
        unique_name = sheet.cell(row_number, 1).value
        if not unique_name:
            continue
        leaf = str(unique_name).rstrip("/").rsplit("/", 1)[-1]
        rows.append(
            {
                "row": row_number,
                "uniqueName": unique_name,
                "name": sheet.cell(row_number, 2).value,
                "leaf": leaf,
            }
        )
    workbook.close()

    unique_leaves = sorted({row["leaf"] for row in rows})
    print(f"Probing {len(rows)} rows / {len(unique_leaves)} unique leaves with {WORKERS} workers", flush=True)
    by_leaf = {}
    with ThreadPoolExecutor(max_workers=WORKERS) as executor:
        futures = {executor.submit(probe_leaf, leaf): leaf for leaf in unique_leaves}
        for completed, future in enumerate(as_completed(futures), 1):
            leaf = futures[future]
            try:
                by_leaf[leaf] = future.result()
            except Exception as error:
                by_leaf[leaf] = {"leaf": leaf, "allimages": [], "search": [], "errors": [{"query": "worker", "error": str(error)}]}
            if completed % 25 == 0 or completed == len(unique_leaves):
                print(f"completed {completed}/{len(unique_leaves)}", flush=True)

    enriched_rows = []
    for row in rows:
        probe = by_leaf[row["leaf"]]
        enriched_rows.append(
            {
                **row,
                "allimages": probe["allimages"],
                "search": probe["search"],
                "errors": probe["errors"],
            }
        )
    payload = {
        "generatedAt": "2026-10-16",
        "source": API,
        "queryPolicy": {
            "allimages": "list=allimages&aiprefix=<leaf>&ailimit=10",
            "search": "list=search&srsearch=<leaf>&srnamespace=0|6&srlimit=10",
            "workers": WORKERS,
            "minimumRequestIntervalSeconds": MIN_REQUEST_INTERVAL,
        },
        "rowCount": len(rows),
        "uniqueLeafCount": len(unique_leaves),
        "rows": enriched_rows,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    print(f"wrote {OUT}", flush=True)


if __name__ == "__main__":
    main()
